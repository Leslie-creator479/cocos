"""학원에서 쓰는 시중 단어장(예: 능률VOCA 고등기본)을 단어게임 페이지로 만든다. (2026-10-17 선생님 결정)

자료실 '단어장' 분류(game/voca-<약칭>)에 들어가고, 시중 교재라 학원생 비밀번호로 잠긴다(lock.mjs가 game/ 아래 모의고사 이름이 아닌 것을 잠금).
단어 데이터는 저장소에 평문으로 두지 않는다(공개 저장소라 교재 내용이 그대로 드러남). 작업 폴더(스크래치패드)에만 둔다.

1) 엑셀(학원 단어 시험지 프로그램 출력, 'rand' 시트: C 단어 · D 뜻 · J '(Day 01) develop')에서 단어 꺼내기
   python3 _tools/voca_game.py extract 시험지1.xlsx [시험지2.xlsx …] -o words.json
   → [{"day":1,"w":"develop","m":"…","p":"동"}, …]. p(품사)는 첫 뜻으로 추정한 값이니 꼭 한 번 훑어 고친다
     (보기에서 품사로 답이 보이지 않게 하는 기준. 명·동·형·부·숙)
2) 이미 올린 페이지에 Day를 더할 때: 잠긴 페이지를 복사해 unlock 한 뒤
   python3 _tools/voca_game.py from-page 풀린페이지.html -o old.json   → 1)과 같은 모양. 두 목록을 합쳐(같은 day는 새것) 3)으로
3) 페이지 만들기
   python3 _tools/voca_game.py build words.json meta.json game/voca-nehs-basic/index.html
   meta.json: {"title":"능률VOCA 고등기본","key":"cocos-wordgame-voca-nehs-basic","chip":"능률VOCA 고등기본 · 정규반 단어 시험",
               "source":"단어 출처: 능률VOCA 고등기본 (NE능률)","group":10,
               "buy":{"label":"정규반 단어 교재: 능률VOCA 고등기본","text":"교보문고에서 사기","url":"https://…"}}
   - 단어장 한 권(Day 01~40 전부)을 한 페이지에 넣고 Day 하나씩 고르게 한다(반마다 두 과·세 과씩 시험 범위가 달라서, 선생님 결정 2026-10-17).
     묶음 버튼은 빨리 고르기용으로 group(10)일씩: Day 01~10, 11~20 …
   - 틀: game/2022-03-g2(모의고사 단어게임, 잠기지 않은 페이지)를 그대로 쓰고 단어장에 맞게 고친다
     (지문 → Day, '지문 속 빈칸' 게임 뺌(예문 없음), 학습지·동의어 링크 없음, 머리말 아래 교재 구매 안내).
   만든 뒤: brand_links.py → seo.py(PDF 버튼) → lock.mjs lock, 자료실은
   python3 _tools/add_library.py game "능률VOCA 고등기본" game/voca-nehs-basic <Day 수> --group "정규반 단어장"   (카드 아래 구매 링크는 넣지 않는다)
   범위 링크: 주소#2 (Day 02), 주소#4-6 (Day 04~06), 주소#4-5.9 (Day 04·05·09)
"""
import json, re, sys, os, argparse

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = os.path.join(ROOT, 'game', '2022-03-g2', 'index.html')
V = re.compile(r'(하다|되다|시키다|내다|지다|주다|두다|리다|기다|우다|추다|르다|치다|히다|이다)$')
A = re.compile(r'(적인|한|운|인|된|진|는|의|있는|없는|은|던)$')


def guess(w, m):
    if ' ' in w.strip():
        return '숙'
    first = re.sub(r'\(.*?\)|\[.*?\]', '', re.split(r'[;,]', m)[0]).strip()
    if A.search(first):
        return '형'
    if V.search(first) or first.endswith('다'):
        return '동'
    if re.search(r'(게|히|아마|매우|거의)$', first):
        return '부'
    return '명'


def opened(f):
    """엑셀을 연다. 가끔 '읽기 전용 암호'(Excel 기본 암호)가 걸린 채로 저장된 파일이 오는데(2026-10-17 워드마스터 Day 13~14),
    msoffcrypto로 풀고 openpyxl이 못 읽는 이름 없는 셀 스타일을 고쳐서 연다."""
    import openpyxl, zipfile, io
    if not zipfile.is_zipfile(f):
        import msoffcrypto  # pip install msoffcrypto-tool --break-system-packages
        o = msoffcrypto.OfficeFile(open(f, 'rb')); o.load_key(password='VelvetSweatshop')
        buf = io.BytesIO(); o.decrypt(buf); buf.seek(0)
        zin = zipfile.ZipFile(buf); fixed = io.BytesIO(); zout = zipfile.ZipFile(fixed, 'w', zipfile.ZIP_DEFLATED); n = [0]
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == 'xl/styles.xml':
                def fx(mm):
                    t = mm.group(0)
                    if ' name=' not in t:
                        n[0] += 1; t = t.replace('cellStyle ', f'cellStyle name="s{n[0]}" ', 1)
                    return t
                d = re.sub(r'<(?:x:)?cellStyle [^>]*>', fx, d.decode('utf8')).encode('utf8')
            zout.writestr(it, d)
        zout.close(); fixed.seek(0); f = fixed
        print('암호 풀어 읽음')
    return openpyxl.load_workbook(f, data_only=True, read_only=True)


def extract(files, out):
    rows, seen = [], set()
    for f in files:
        wb = opened(f); ws = wb['rand']
        recs = [r for r in ws.iter_rows(min_row=2, values_only=True) if len(r) >= 10 and r[2]]
        # 'Day 표시'를 끄고 뽑은 파일(2026-10-17 워드마스터)은 J열에 (Day 01)이 없다 → 머리말 '출제범위 … Day 01 / Day 02'의
        # Day들로 순번(A열) 순서대로 똑같이 나눈다(Day마다 단어 수가 같은 책일 때만. 다르면 멈춘다)
        if not any(re.match(r'\(Day \d+\)', str(r[9] or ''), re.I) for r in recs):  # 해커스는 'DAY 01'(대문자)
            head = next((str(c) for row in wb['print'].iter_rows(max_row=5, values_only=True) for c in row if c and '출제범위' in str(c)), '')
            ds = [int(x) for x in re.findall(r'Day (\d+)', head, re.I)]
            assert ds and len(recs) % len(ds) == 0, f'Day 표시도 없고 Day별로 똑같이 나눌 수 없음: {f}'
            per = len(recs) // len(ds); recs.sort(key=lambda r: r[0])
            recs = [tuple(r[:9]) + ((f'(Day {ds[i // per]:02d}) ' if i % per == 0 else '') + str(r[2]),) for i, r in enumerate(recs)]
            print(f'Day 표시 없음 → {ds} 를 {per}개씩 순서대로 나눔')
        day = None
        for r in recs:
            mm = re.match(r'\(Day (\d+)\)', str(r[9] or ''), re.I)
            if mm:
                day = int(mm.group(1))
            w, m = str(r[2]).strip(), str(r[3]).strip()
            if (day, w) in seen:
                continue
            seen.add((day, w)); rows.append(dict(day=day, w=w, m=m, p=guess(w, m)))
    json.dump(rows, open(out, 'w', encoding='utf8'), ensure_ascii=False, indent=0)
    from collections import Counter
    print(len(rows), '단어 · Day별', dict(sorted(Counter(r['day'] for r in rows).items())), '· 품사(추정)', dict(Counter(r['p'] for r in rows)))


def from_page(page, out):
    s = open(page, encoding='utf8').read()
    a = s.index('var DATA = ') + len('var DATA = ')
    D, _ = json.JSONDecoder().raw_decode(s[a:])
    days = [int(p['k'].lstrip('d')) for p in D['P']]
    rows = [dict(day=days[x[3]], w=x[0], m=x[2], p=x[1]) for x in D['W']]
    json.dump(rows, open(out, 'w', encoding='utf8'), ensure_ascii=False, indent=0)
    print(len(rows), '단어 · Day', days)


PATCH = [
    ("  function speak(t){ if (!synth) return; try { synth.cancel(); var u = new SpeechSynthesisUtterance(t); u.lang = 'en-US'; u.rate = 0.9; if (voice) u.voice = voice; synth.speak(u); } catch (e) {} }", "  function tts(t){ if (!synth) return; try { synth.cancel(); var u = new SpeechSynthesisUtterance(t); u.lang = 'en-US'; u.rate = 0.9; if (voice) u.voice = voice; synth.speak(u); } catch (e) {} }\n  /* 발음은 미리 녹음한 MP3(assets/say/<단어>.mp3)를 먼저 튼다. 카카오톡 안 브라우저는 기기 음성이 없어서(2026-10-17) */\n  var SAY = DATA.say != null ? DATA.say : '../../assets/say/', player = null;\n  function slug(w){ return String(w).toLowerCase().replace(/[^a-z0-9]+/g, '-').replace(/^-+|-+$/g, ''); }\n  function speak(t){ if (SAY) { try { if (!player) player = new Audio(); player.pause(); player.src = SAY + slug(t) + '.mp3'; var pr = player.play(); if (pr && pr.catch) pr.catch(function(){ tts(t); }); return; } catch (e) {} } tts(t); }"),
    ('  function sayBtn(w){ if (!synth) return null;', '  function sayBtn(w){ if (!synth && !SAY) return null;'),
    ("return '전체 ' + P.length + '개 지문';", "return '전체 Day ' + P.length + '개';"),
    ("return gs.length + '개 지문';", "return 'Day ' + gs.length + '개';"),
    ("공부할 지문 번호를 눌러 고르세요.", "공부할 Day를 눌러 고르세요."),
    ("지문 번호를 눌러 골라요.", "Day를 눌러 골라요."),
    ("공부할 지문을 먼저", "공부할 Day를 먼저"),
    ("'aria-label': '지문 번호 고르기'", "'aria-label': 'Day 고르기'"),
    ("온라인 학습지의 단어 정리를 지문별로 모았어요.", "단어장 단어를 Day별로 모았어요."),
    ("guide, tip, list, go", "tip, list, go"),
    ("var ORDER = ['match', 'ta', 'fall', 'spell', 'blank', 'review'];", "var ORDER = ['match', 'ta', 'fall', 'spell', 'review'];"),
    ("lab(p) + ' · ' + p.topic, h('small'", "lab(p) + (p.topic ? ' · ' + p.topic : ''), h('small'"),
    ("    var e = document.createElement('a'); e.href = '#academy'; e.textContent = '코코스학원'; nav.append(e);\n",
     "    var e = document.createElement('a'); e.href = '#academy'; e.textContent = '코코스학원'; nav.append(e);\n"
     "    if (M.buy) { var bx = document.createElement('p'); bx.className = 'buybox'; var bb = document.createElement('b'); bb.textContent = M.buy.label; bx.append(bb);"
     " var ba = document.createElement('a'); ba.href = M.buy.url; ba.target = '_blank'; ba.rel = 'noopener'; ba.textContent = (M.buy.text || '구매하기') + ' ↗'; bx.append(ba); nav.parentNode.insertBefore(bx, nav); }\n"),
    ("</style>", ".buybox{display:flex;flex-wrap:wrap;align-items:center;gap:6px 12px;margin:12px 0 10px;padding:10px 14px;border-radius:12px;background:#FFF7D6;border:1px solid #F2D45C;font-size:15px}"
     ".buybox b{font-weight:700;color:#5A4500}.buybox a{color:#065A88;font-weight:700;text-decoration:none;border-bottom:1px solid currentColor}</style>"),
]


def build(words_f, meta_f, out):
    W = json.load(open(words_f, encoding='utf8')); M = json.load(open(meta_f, encoding='utf8'))
    days = sorted({x['day'] for x in W})
    P = [dict(no=f'Day {d:02d}', k=str(d), lab=f'Day {d:02d}', topic='', u='', s=[]) for d in days]
    GW = [[x['w'], x['p'], x['m'], days.index(x['day'])] for x in sorted(W, key=lambda x: x['day'])]
    g = M.get('group', 10)
    groups, d0 = [], 1
    while d0 <= max(days):
        d1 = d0 + g - 1
        ks = [str(d) for d in range(d0, d1 + 1) if d in days]
        if ks:
            groups.append(dict(label=f'Day {d0:02d}~{d1:02d}', keys=ks))
        d0 = d1 + 1
    rng = f'Day {days[0]:02d}~{days[-1]:02d}'
    meta = dict(h1=f"{M['title']} 단어게임", chips=[M['chip'], f'{rng} · 단어 {len(GW)}개', '게임 4가지'],
                gist='정규반 단어 시험 범위를 골라 단어를 먼저 익히고, 네 가지 게임으로 확인해 보세요. 틀린 단어는 ‘오답 다시 풀기’에 모여요.',
                links=[], source=M['source'])
    if M.get('buy'):
        meta['buy'] = M['buy']
    s = open(TPL, encoding='utf8').read()
    if '<body>' in s[:5000]:
        s = s.split('<body>', 1)[1].rsplit('</body></html>', 1)[0].lstrip('\n')
    s = re.sub(r'\n*<script src="[./]*assets/print-game\.js"></script>\n*', '\n', s)
    a = s.index('/* ==== DATA:START ==== */'); b = s.index('/* ==== DATA:END ==== */')
    head = s[a:s.index('var DATA = ', a)]
    data = dict(unit='Day', say='../../assets/say/', key=M['key'], label=M['title'], syn='', meta=meta, groups=groups, P=P, W=GW)
    s = s[:a] + head + 'var DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n' + s[b:]
    s = re.sub(r'<title>[^<]*</title>', f"<title>[코코스학원] {M['title']} 단어게임</title>", s, count=1)
    for x, y in PATCH:
        if x not in s and all(l in s for l in y.split('\n') if l.strip() and not l.strip().startswith('/*')):
            continue  # 틀(game/2022-03-g2)에 이미 들어간 고침(예: 녹음 발음, 2026-10-17 이후 모든 단어게임에 반영됨)
        assert x in s, '틀이 바뀜: ' + x[:50]
        s = s.replace(x, y)
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    tmp = out + '.tmp.html'
    open(tmp, 'w', encoding='utf8').write(s)
    os.system(f'python3 {ROOT}/_tools/wrap.py {tmp} {out} >/dev/null'); os.remove(tmp)
    if M.get('mid', bool(re.search('중등|중학', M['title']))):  # 중등·중학 단어장은 중등 테마(원장님 결정 2026-10-25, meta.json "mid": false로 끌 수 있음)
        os.system(f'python3 {ROOT}/_tools/mid_sheet/theme_page.py game {out} >/dev/null')
    print(out, '·', rng, '· 단어', len(GW), '· 묶음', [x['label'] for x in groups])


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest='cmd', required=True)
    e = sp.add_parser('extract'); e.add_argument('files', nargs='+'); e.add_argument('-o', required=True)
    f = sp.add_parser('from-page'); f.add_argument('page'); f.add_argument('-o', required=True)
    b = sp.add_parser('build'); b.add_argument('words'); b.add_argument('meta'); b.add_argument('out')
    a = ap.parse_args()
    if a.cmd == 'extract': extract(a.files, a.o)
    elif a.cmd == 'from-page': from_page(a.page, a.o)
    else: build(a.words, a.meta, a.out)
